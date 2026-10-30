"""Tests for the .docx form export (Phase 5)."""

from __future__ import annotations

import base64
import io

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT

from app.services.form_docx import (
    _as_str_list,
    build_form_document,
    render_form_docx,
    sanitize_filename,
)

DOC = {
    "version": 1,
    "creation_mode": "structured",
    "title": "CET Form",
    "subtitle": "College Entrance Test",
    "instructions": "Fill in all blanks.",
    "letterhead": {
        "institutionName": "CCS",
        "addressLines": ["123 Main St"],
        "contactLine": "contact@example.com",
        "formNumber": "F-001",
    },
    "page": {
        "size": "A4",
        "orientation": "portrait",
        "margins": {"top": 14, "right": 16, "bottom": 14, "left": 16},
    },
    "sections": [
        {
            "id": "s1",
            "heading": "Student Info",
            "blocks": [
                {"type": "field", "id": "b1", "field_key": "full_name", "label": "Full Name", "ui_component": "text_input", "required": True, "options": None},
                {"type": "field", "id": "b2", "field_key": "sex", "label": "Sex", "ui_component": "radio_group", "required": False, "options": [{"value": "m", "label": "Male"}, {"value": "f", "label": "Female"}]},
                {"type": "paragraph", "id": "p1", "text": "Please print clearly."},
                {"type": "table", "id": "t1", "columns": 2, "headerRow": True, "rows": [["Subject", "Grade"], ["Math", "95"]]},
                {"type": "signature", "id": "sig1", "label": "Signature over Printed Name"},
            ],
        }
    ],
    "signatureLabel": "Signature over Printed Name",
    "footer": "Keep this form.",
}

# Minimal valid 1x1 transparent PNG.
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
)


def _paragraph_texts(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs]


def _all_text(doc: Document) -> list[str]:
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    return texts


def _table_with(doc: Document, needle: str) -> object | None:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if needle in cell.text:
                    return table
    return None


def test_render_form_docx_returns_parseable_docx_bytes() -> None:
    data = render_form_docx(DOC)

    assert isinstance(data, bytes)
    assert data[:2] == b"PK"  # .docx is a ZIP archive
    Document(io.BytesIO(data))  # must parse without error


def test_build_form_document_renders_letterhead_title_and_sections() -> None:
    doc = build_form_document(DOC)
    texts = _all_text(doc)

    def has(needle: str) -> bool:
        return any(needle in t for t in texts)

    # Letterhead (uppercase institution, address, contact, form number).
    assert has("CCS")
    assert has("123 Main St")
    assert has("contact@example.com")
    assert has("Form No.")
    assert has("F-001")
    # Title is centred + uppercase.
    assert has("CET FORM")
    assert has("College Entrance Test")
    assert has("Fill in all blanks.")
    # Field labels (uppercase, required suffix preserved).
    assert has("FULL NAME *")
    assert has("SEX")
    assert has("Please print clearly.")
    # Footer uppercase.
    assert has("KEEP THIS FORM.")


def test_build_form_document_uppercases_field_labels() -> None:
    doc = build_form_document(DOC)
    texts = _all_text(doc)

    assert any("FULL NAME" in t for t in texts)


def test_build_form_document_uses_serif_for_title() -> None:
    doc = build_form_document(DOC)
    title_paragraphs = [p for p in doc.paragraphs if "CET FORM" in p.text]

    assert title_paragraphs
    title_run = title_paragraphs[0].runs[0]
    assert title_run.font.name == "Cambria"
    assert title_run.bold is True


def test_build_form_document_uses_radio_symbol_for_radio_group() -> None:
    doc = build_form_document(DOC)
    texts = _all_text(doc)

    assert any("\u25cb" in t for t in texts)  # hollow circle for radio


def test_build_form_document_renders_inline_label_colon_and_line() -> None:
    doc = build_form_document(DOC)
    texts = _all_text(doc)

    # Label with a colon on the same row as the write-on line.
    assert any("FULL NAME *:" in t for t in texts)
    # The write-on line is rendered with non-breaking spaces (so Word underlines it).
    assert any("\u00A0" in t for t in texts)
    # Choice fields list their options inline after the label.
    assert any(("SEX:" in t) and ("\u25cb" in t) for t in texts)


def test_build_form_document_renders_field_options() -> None:
    doc = build_form_document(DOC)
    texts = _all_text(doc)

    assert any("Male" in t for t in texts)
    assert any("Female" in t for t in texts)


def test_build_form_document_renders_table() -> None:
    doc = build_form_document(DOC)

    table = _table_with(doc, "Subject")
    assert table is not None
    assert len(table.rows) == 2
    assert table.rows[0].cells[0].text == "Subject"
    assert table.rows[1].cells[1].text == "95"


def test_build_form_document_pairs_half_width_fields() -> None:
    doc = build_form_document(DOC)

    # "FULL NAME" and "SEX" are both columnSpan=1, so they share a 2-column row.
    grid = _table_with(doc, "FULL NAME")
    assert grid is not None
    assert len(grid.columns) == 2
    cell_texts = [cell.text for cell in grid.rows[0].cells]
    assert any("FULL NAME" in t for t in cell_texts)
    assert any("SEX" in t for t in cell_texts)


def test_build_form_document_embeds_logo_when_provided() -> None:
    doc = build_form_document(DOC, logo_bytes=PNG_1PX)

    assert len(doc.inline_shapes) == 1


def test_build_form_document_skips_logo_when_absent() -> None:
    doc = build_form_document(DOC, logo_bytes=None)

    assert len(doc.inline_shapes) == 0


def test_build_form_document_applies_page_settings() -> None:
    doc = build_form_document(
        {
            "version": 1,
            "title": "Landscape",
            "page": {
                "size": "Letter",
                "orientation": "landscape",
                "margins": {"top": 10, "right": 12, "bottom": 10, "left": 12},
            },
        }
    )

    section = doc.sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width.mm == pytest.approx(279.4, abs=0.5)
    assert section.page_height.mm == pytest.approx(215.9, abs=0.5)
    assert section.left_margin.mm == pytest.approx(12, abs=0.1)
    assert section.top_margin.mm == pytest.approx(10, abs=0.1)


def test_build_form_document_legacy_fields_fallback() -> None:
    legacy = {
        "version": 1,
        "creation_mode": "structured",
        "title": "Legacy",
        "sections": [
            {
                "id": "s1",
                "heading": "Section",
                "fields": [
                    {"id": "f1", "field_key": "a", "label": "Alpha", "ui_component": "text_input", "required": False},
                ],
            }
        ],
    }
    doc = build_form_document(legacy)
    assert any("ALPHA" in t for t in _all_text(doc))


def test_build_form_document_empty_document() -> None:
    doc = build_form_document({"version": 1, "title": "Empty"})
    assert "EMPTY" in _paragraph_texts(doc)


def test_build_form_document_empty_options_renders_write_in_line() -> None:
    doc = build_form_document(
        {
            "version": 1,
            "title": "T",
            "sections": [
                {
                    "id": "s1",
                    "blocks": [
                        {"type": "field", "id": "f1", "field_key": "pick", "label": "Pick", "ui_component": "radio_group", "options": None},
                    ],
                }
            ],
        }
    )
    texts = _all_text(doc)

    # Empty options fall back to a write-on line (no option symbols).
    assert not any("\u25cb" in t for t in texts)
    assert any("\u00A0" in t for t in texts)


def test_build_form_document_dropdown_renders_write_in_line_not_checkbox() -> None:
    doc = build_form_document(
        {
            "version": 1,
            "title": "T",
            "sections": [
                {
                    "id": "s1",
                    "blocks": [
                        {"type": "field", "id": "f1", "field_key": "course", "label": "Course", "ui_component": "dropdown", "options": [{"value": "bsit", "label": "BSIT"}]},
                    ],
                }
            ],
        }
    )
    texts = _all_text(doc)

    # Dropdown renders a write-on line, not a checkbox list.
    assert any("COURSE:" in t for t in texts)
    assert any("\u00A0" in t for t in texts)
    assert not any("\u2610" in t for t in texts)


def test_build_form_document_date_picker_renders_format_hint() -> None:
    doc = build_form_document(
        {
            "version": 1,
            "title": "T",
            "sections": [
                {
                    "id": "s1",
                    "blocks": [
                        {"type": "field", "id": "f1", "field_key": "dob", "label": "Birth Date", "ui_component": "date_picker"},
                    ],
                }
            ],
        }
    )
    texts = _all_text(doc)

    assert any("BIRTH DATE: MM / DD / YYYY" in t for t in texts)


def test_sanitize_filename_strips_unsafe_characters() -> None:
    assert sanitize_filename("CET Form") == "CET Form"
    assert sanitize_filename('A/B\\C"D') == "A-B-CD"
    assert sanitize_filename("  Résumé Form  ") == "Rsum Form"
    assert sanitize_filename(None) == "form"
    assert sanitize_filename("   ") == "form"


def _doc_with_field(field: dict) -> Document:
    return build_form_document(
        {"version": 1, "title": "T", "sections": [{"id": "s1", "blocks": [field]}]}
    )


def test_build_form_document_full_width_field_not_paired() -> None:
    doc = _doc_with_field(
        {"type": "field", "id": "f1", "field_key": "full", "label": "Full Field", "ui_component": "text_input", "columnSpan": 2}
    )
    # A full-width (columnSpan >= 2) field renders directly in the body, not in a table cell.
    assert any("FULL FIELD:" in t for t in _paragraph_texts(doc))


def test_build_form_document_checkbox_uses_ballot_box_symbol() -> None:
    doc = _doc_with_field(
        {"type": "field", "id": "f1", "field_key": "subjects", "label": "Subjects", "ui_component": "checkbox_group", "options": [{"value": "math", "label": "Math"}]}
    )
    texts = _all_text(doc)

    assert any("\u2610" in t for t in texts)  # ballot box for checkbox
    assert not any("\u25cb" in t for t in texts)  # no radio circle


def test_build_form_document_empty_label_falls_back_to_field_key() -> None:
    doc = _doc_with_field(
        {"type": "field", "id": "f1", "field_key": "student_name", "label": "", "ui_component": "text_input"}
    )
    assert any("STUDENT_NAME:" in t for t in _all_text(doc))


def test_build_form_document_uses_branding_when_letterhead_absent() -> None:
    doc = build_form_document(
        {"version": 1, "title": "T"},
        branding={"institution_name": "Test University", "address_lines": ["1 Test St"], "contact_line": "test@example.com"},
    )
    texts = _all_text(doc)

    assert any("TEST UNIVERSITY" in t for t in texts)
    assert any("1 Test St" in t for t in texts)
    assert any("test@example.com" in t for t in texts)


def test_as_str_list_handles_string_list_and_empty() -> None:
    assert _as_str_list("hello") == ["hello"]
    assert _as_str_list(["a", "b"]) == ["a", "b"]
    assert _as_str_list(None) == []
    assert _as_str_list("") == []


def test_build_form_document_skips_unresolved_field() -> None:
    doc = build_form_document(
        {
            "version": 1,
            "title": "T",
            "sections": [
                {
                    "id": "s1",
                    "blocks": [
                        {"type": "unresolved_field", "id": "u1", "field_key": "invented", "label": "Invented"},
                        {"type": "field", "id": "f1", "field_key": "real", "label": "Real", "ui_component": "text_input"},
                    ],
                }
            ],
        }
    )
    texts = _all_text(doc)

    assert any("REAL:" in t for t in texts)
    assert not any("INVENTED" in t for t in texts)


def test_build_form_document_a4_portrait_page_settings() -> None:
    doc = build_form_document(
        {"version": 1, "title": "T", "page": {"size": "A4", "orientation": "portrait", "margins": {"top": 10, "right": 12, "bottom": 10, "left": 12}}}
    )
    section = doc.sections[0]

    assert section.orientation == WD_ORIENT.PORTRAIT
    assert section.page_width.mm == pytest.approx(210, abs=0.5)
    assert section.page_height.mm == pytest.approx(297, abs=0.5)
    assert section.top_margin.mm == pytest.approx(10, abs=0.1)
    assert section.left_margin.mm == pytest.approx(12, abs=0.1)


def test_build_form_document_field_style_applied() -> None:
    doc = _doc_with_field(
        {"type": "field", "id": "f1", "field_key": "styled", "label": "Styled", "ui_component": "text_input", "style": {"fontSize": 16, "italic": True}}
    )
    styled_para = next((p for p in doc.paragraphs if "STYLED:" in p.text), None)

    assert styled_para is not None
    assert any(r.italic for r in styled_para.runs)
    assert any(r.font.size is not None and r.font.size.pt == 16 for r in styled_para.runs)


def test_build_form_document_signature_custom_label() -> None:
    doc = _doc_with_field({"type": "signature", "id": "sig1", "label": "Applicant Signature"})
    assert any("APPLICANT SIGNATURE" in t for t in _all_text(doc))


# ── export_form_docx endpoint ─────────────────────────────────────────────


@pytest.mark.asyncio
async def test_export_form_docx_returns_404_for_missing_form() -> None:
    from unittest.mock import AsyncMock
    from uuid import uuid4

    from fastapi import HTTPException

    from app.routers.admin.forms import export_form_docx

    db = AsyncMock()
    db.get = AsyncMock(return_value=None)

    with pytest.raises(HTTPException) as exc:
        await export_form_docx(uuid4(), {}, db)

    assert exc.value.status_code == 404


@pytest.mark.asyncio
async def test_export_form_docx_returns_docx_with_headers() -> None:
    from types import SimpleNamespace
    from unittest.mock import AsyncMock
    from uuid import uuid4

    from app.routers.admin.forms import export_form_docx

    form = SimpleNamespace(id=uuid4(), name="CET Form", document_json=DOC)
    db = AsyncMock()
    # First get() -> the form; second get() -> BrandingSettings (None, no logo).
    db.get = AsyncMock(side_effect=[form, None])

    resp = await export_form_docx(form.id, {}, db)

    assert resp.status_code == 200
    assert resp.media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert resp.body[:2] == b"PK"
    assert 'filename="CET Form.docx"' in resp.headers.get("content-disposition", "")
