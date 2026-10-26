"""Render a structured form ``document_json`` to a Word ``.docx`` for download.

Phase 5 ``.docx`` export. The output mirrors the printed form's typography and
layout as closely as Word allows: a serif letterhead (logo + institution + form
number), a serif centred uppercase title, sans body text, numbered section
headings, a two-column field grid (``columnSpan``-aware), paragraphs, tables,
signatures, and a footer. Fields render as an uppercase label plus an underlined
blank line (or an option list for choice widgets). No AI, no external service.
"""

from __future__ import annotations

import io
import logging
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

logger = logging.getLogger(__name__)

_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_CHOICE_COMPONENTS = {"radio_group", "checkbox_group"}

# Preview font families (font-serif → Cambria; body uses Inter → Calibri fallback;
# font-mono → Courier New).
_SERIF = "Cambria"
_SANS = "Calibri"
_MONO = "Courier New"

_DEFAULT_INSTITUTION = "College of Computer Studies"
_DEFAULT_ADDRESS = ["Campus Address, City, Province"]
_DEFAULT_CONTACT = "registrar@example.edu \u00b7 (000) 000-0000"
_DEFAULT_FORM_NUMBER = "ADS-01"


def _font(run: Any, name: str | None = None, size: float | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    if name:
        run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _set_bottom_border(paragraph: Any, size: int = 6, color: str = "1f2937") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _apply_style(paragraph: Any, style: dict[str, Any] | None) -> None:
    if not style:
        return
    align = _ALIGN_MAP.get(style.get("align"))
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        run.bold = bool(style.get("bold", False)) or run.bold
        run.italic = bool(style.get("italic", False)) or run.italic
        size = style.get("fontSize")
        if size:
            run.font.size = Pt(size)


def _add_write_in_line(container: Any) -> None:
    """Draw a standalone full-width write-on line (bottom border), used by the
    signature + date block."""
    paragraph = container.add_paragraph()
    run = paragraph.add_run("\u00A0")
    _font(run, size=12)
    _set_bottom_border(paragraph, size=4, color="94a3b8")
    paragraph.paragraph_format.space_after = Pt(10)


def _option_symbol(ui: str) -> str:
    if ui == "radio_group":
        return "\u25cb"  # hollow circle
    return "\u2610"  # ballot box (checkbox)


def _render_field(container: Any, block: dict[str, Any]) -> None:
    label = block.get("label") or block.get("field_key") or ""
    required = bool(block.get("required", False))
    ui = block.get("ui_component", "text_input")

    label_text = label.upper() + (" *" if required else "")
    paragraph = container.add_paragraph()
    label_run = paragraph.add_run(f"{label_text}: ")
    _font(label_run, name=_SANS, size=8, bold=True)

    options = block.get("options") or []
    if ui in _CHOICE_COMPONENTS and options:
        symbol = _option_symbol(ui)
        first = True
        for opt in options:
            opt_label = opt.get("label", opt.get("value", "")) if isinstance(opt, dict) else str(opt)
            sep = "" if first else "   "
            opt_run = paragraph.add_run(f"{sep}{symbol} {opt_label}")
            _font(opt_run, name=_SANS, size=8)
            first = False
    elif ui == "date_picker":
        # Format hint with an underline (the write-on line) for the date.
        hint_run = paragraph.add_run("MM / DD / YYYY")
        hint_run.underline = True
        _font(hint_run, name=_SANS, size=8)
        hint_run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)  # slate-400
    else:
        # Inline write-on line: underlined non-breaking spaces (regular trailing
        # spaces are not underlined by Word).
        line_run = paragraph.add_run("\u00A0" * 35)
        line_run.underline = True

    _apply_style(paragraph, block.get("style"))


def _render_field_row(doc: Document, fields: list[dict[str, Any]]) -> None:
    if len(fields) == 1:
        _render_field(doc, fields[0])
        return

    table = doc.add_table(rows=1, cols=len(fields))
    table.autofit = True
    for idx, field in enumerate(fields):
        cell = table.cell(0, idx)
        cell.width = Inches(3.2)
        _render_field(cell, field)


def _render_paragraph(doc: Document, block: dict[str, Any]) -> None:
    paragraph = doc.add_paragraph(block.get("text", ""))
    for run in paragraph.runs:
        _font(run, name=_SANS, size=8)
    _apply_style(paragraph, block.get("style"))


def _render_table(doc: Document, block: dict[str, Any]) -> None:
    rows: list[list[str]] = block.get("rows") or []
    columns = max(int(block.get("columns", 2) or 2), 1)
    header_row = bool(block.get("headerRow", False))

    if not rows:
        return

    table = doc.add_table(rows=0, cols=columns)
    table.style = "Table Grid"

    for row_index, row in enumerate(rows):
        cells = row if isinstance(row, list) else [str(row)]
        cells = (cells + [""] * columns)[:columns]
        table_row = table.add_row().cells
        for col_index, text in enumerate(cells):
            table_row[col_index].text = str(text)
            for p in table_row[col_index].paragraphs:
                for r in p.runs:
                    _font(r, name=_SANS, size=8)
                    if header_row and row_index == 0:
                        r.bold = True

    doc.add_paragraph()


def _render_signature_block(doc: Document, block: dict[str, Any]) -> None:
    label = block.get("label", "Signature over Printed Name")
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True

    left = table.cell(0, 0)
    _add_write_in_line(left)
    p = left.add_paragraph(label.upper())
    for run in p.runs:
        _font(run, name=_SANS, size=7)

    right = table.cell(0, 1)
    _add_write_in_line(right)
    p = right.add_paragraph("DATE")
    for run in p.runs:
        _font(run, name=_SANS, size=7)


def _render_block(doc: Document, block: dict[str, Any]) -> None:
    block_type = block.get("type")
    if block_type == "field":
        _render_field(doc, block)
    elif block_type == "paragraph":
        _render_paragraph(doc, block)
    elif block_type == "table":
        _render_table(doc, block)
    elif block_type == "signature":
        _render_signature_block(doc, block)
    # unresolved_field is intentionally skipped (never persisted).


def _render_section_heading(doc: Document, index: int, heading: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(f"{index + 1:02d}. {heading.upper()}")
    _font(run, name=_SANS, size=9, bold=True)
    _set_bottom_border(paragraph, size=4, color="cbd5e1")


def _render_letterhead(
    doc: Document,
    institution: str,
    address_lines: list[str],
    contact: str,
    form_number: str,
    logo_bytes: bytes | None,
) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Inches(1.0)
    table.columns[1].width = Inches(4.0)
    table.columns[2].width = Inches(1.2)

    logo_cell = table.cell(0, 0)
    if logo_bytes:
        try:
            logo_cell.paragraphs[0].add_run().add_picture(io.BytesIO(logo_bytes), width=Inches(0.9))
        except Exception:  # noqa: BLE001 - image bytes may be corrupt/unrecognized in several ways
            logger.warning("Failed to embed branding logo in .docx export; exporting without logo")

    center_cell = table.cell(0, 1)
    name_para = center_cell.paragraphs[0]
    name_run = name_para.add_run(institution.upper())
    _font(name_run, name=_SERIF, size=13, bold=True)
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in address_lines:
        p = center_cell.add_paragraph(str(line))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            _font(run, name=_SANS, size=8)
    contact_para = center_cell.add_paragraph(contact)
    contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in contact_para.runs:
        _font(run, name=_SANS, size=8)

    number_cell = table.cell(0, 2)
    label_para = number_cell.paragraphs[0]
    label_run = label_para.add_run("Form No.")
    _font(label_run, name=_SANS, size=7, bold=True)
    label_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    num_para = number_cell.add_paragraph(form_number)
    num_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in num_para.runs:
        _font(run, name=_MONO, size=8)

    rule = doc.add_paragraph()
    _set_bottom_border(rule, size=6, color="1f2937")


def _apply_page_settings(doc: Document, page: dict[str, Any] | None) -> None:
    if not page:
        return
    section = doc.sections[0]
    size = page.get("size", "A4")
    orientation = page.get("orientation", "portrait")

    if size == "Letter":
        width, height = Mm(215.9), Mm(279.4)
    else:
        width, height = Mm(210), Mm(297)
    if orientation == "landscape":
        width, height = height, width
        section.orientation = WD_ORIENT.LANDSCAPE
    else:
        section.orientation = WD_ORIENT.PORTRAIT

    section.page_width = width
    section.page_height = height

    margins = page.get("margins") or {}
    section.top_margin = Mm(margins.get("top", 14))
    section.bottom_margin = Mm(margins.get("bottom", 14))
    section.left_margin = Mm(margins.get("left", 16))
    section.right_margin = Mm(margins.get("right", 16))


def _as_str_list(value: Any) -> list[str]:
    if not value:
        return []
    if isinstance(value, list):
        return [str(v) for v in value]
    return [str(value)]


def build_form_document(
    document_json: dict[str, Any],
    logo_bytes: bytes | None = None,
    branding: dict[str, Any] | None = None,
) -> Document:
    """Build a python-docx ``Document`` from a v1 structured form document."""
    doc = Document()
    _apply_page_settings(doc, document_json.get("page"))

    letterhead = document_json.get("letterhead") or {}
    branding = branding or {}

    institution = (
        letterhead.get("institutionName")
        or branding.get("institution_name")
        or _DEFAULT_INSTITUTION
    )
    address_lines = _as_str_list(letterhead.get("addressLines")) or _as_str_list(
        branding.get("address_lines")
    ) or _DEFAULT_ADDRESS
    contact = letterhead.get("contactLine") or branding.get("contact_line") or _DEFAULT_CONTACT
    form_number = letterhead.get("formNumber") or _DEFAULT_FORM_NUMBER

    _render_letterhead(doc, institution, address_lines, contact, form_number, logo_bytes)

    title = document_json.get("title") or "Untitled Form"
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title.upper())
    _font(title_run, name=_SERIF, size=20, bold=True)

    subtitle = document_json.get("subtitle")
    if subtitle:
        sub_para = doc.add_paragraph(subtitle)
        sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in sub_para.runs:
            _font(run, name=_SANS, size=9)

    instructions = document_json.get("instructions")
    if instructions:
        instr_para = doc.add_paragraph(instructions)
        for run in instr_para.runs:
            _font(run, name=_SANS, size=8, italic=True)

    sections = document_json.get("sections") or []
    for index, section in enumerate(sections):
        if not isinstance(section, dict):
            continue
        heading = section.get("heading")
        if heading:
            _render_section_heading(doc, index, heading)

        blocks = section.get("blocks") or []
        if blocks:
            i = 0
            while i < len(blocks):
                block = blocks[i]
                if (
                    isinstance(block, dict)
                    and block.get("type") == "field"
                    and (block.get("columnSpan", 1) or 1) < 2
                ):
                    pair = [block]
                    nxt = blocks[i + 1] if i + 1 < len(blocks) else None
                    if (
                        isinstance(nxt, dict)
                        and nxt.get("type") == "field"
                        and (nxt.get("columnSpan", 1) or 1) < 2
                    ):
                        pair.append(nxt)
                        i += 1
                    _render_field_row(doc, pair)
                elif isinstance(block, dict):
                    _render_block(doc, block)
                i += 1
        else:
            # Legacy sections[].fields[] (read-only fallback).
            for spec in section.get("fields") or []:
                if isinstance(spec, dict):
                    _render_field(doc, spec)

    signature_label = document_json.get("signatureLabel") or "Signature over Printed Name"
    _render_signature_block(doc, {"label": signature_label})

    footer = document_json.get("footer")
    if footer:
        footer_para = doc.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer_para.add_run(footer.upper())
        _font(footer_run, name=_SANS, size=7)

    return doc


def render_form_docx(
    document_json: dict[str, Any],
    logo_bytes: bytes | None = None,
    branding: dict[str, Any] | None = None,
) -> bytes:
    """Render a form document to ``.docx`` bytes."""
    doc = build_form_document(document_json, logo_bytes=logo_bytes, branding=branding)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def sanitize_filename(name: str | None) -> str:
    """Return a safe download filename for a form name.

    Strips path separators, quotes, and control characters that would break the
    ``Content-Disposition`` header, keeps printable ASCII, and falls back to
    ``"form"`` when nothing usable remains.
    """
    cleaned = (name or "").strip()
    cleaned = cleaned.replace("/", "-").replace("\\", "-")
    cleaned = cleaned.replace('"', "").replace("\r", " ").replace("\n", " ").replace("\t", " ")
    cleaned = "".join(c if 32 <= ord(c) < 127 else "" for c in cleaned)
    return cleaned.strip(" .") or "form"
